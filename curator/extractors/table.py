"""Excel sheets and CSV: one location per row (r-<sheet>-<row>), tables stay tables."""
from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl

from curator.extractors.base import ExtractedDoc, ExtractionError, Section
from curator.extractors.text import decode
from curator.util import slugify


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\n", " ").replace("|", "/").strip()


def _sheet_sections(sheet_name: str, rows: list[tuple[int, list[str]]]) -> list[Section]:
    rows = [(i, cells) for i, cells in rows if any(cells)]
    if not rows:
        return []
    # useful width: up to the last non-empty column
    width = max(max((j for j, c in enumerate(cells) if c), default=0) for _, cells in rows) + 1
    header_i, header = rows[0][0], rows[0][1][:width]
    header_txt = "| " + " | ".join(header) + " |"
    sslug = slugify(sheet_name)[:30]
    out = []
    for i, cells in rows:
        line = "| " + " | ".join(cells[:width]) + " |"
        out.append(Section(f"r-{sslug}-{i}", f"sheet «{sheet_name}», row {i}", line,
                           group=sheet_name, header="" if i == header_i else header_txt))
    return out


def extract_table(path: Path, **_) -> ExtractedDoc:
    sections: list[Section] = []
    meta: dict = {}
    if path.suffix.lower() == ".csv":
        text, enc = decode(path.read_bytes())
        sample = text[:2000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        rows = [(i, [_fmt(c) for c in r]) for i, r in enumerate(csv.reader(io.StringIO(text), dialect), start=1)]
        sections = _sheet_sections(path.stem, rows)
        meta = {"encoding": enc, "sheets": [path.stem]}
    else:
        try:
            wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        except Exception as e:
            raise ExtractionError(f"unreadable spreadsheet: {e}") from e
        meta["sheets"] = wb.sheetnames
        for ws in wb.worksheets:
            rows = [(i, [_fmt(c) for c in r]) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
            sections.extend(_sheet_sections(ws.title, rows))
        wb.close()
    return ExtractedDoc("table", sections, meta)
