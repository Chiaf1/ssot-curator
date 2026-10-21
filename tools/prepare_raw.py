"""Prepares the raw/ folder of a test instance from a folder of samples, REDACTING sensitive data
(names, customer, passwords, emails) in the original file formats.

It is not part of the curator (which never writes to raw/): it is an assisted human step, to be run
before putting real data in an instance. The samples are never modified.

Usage:
    python tools/prepare_raw.py <config.yaml>            # copy, redact, verify
    python tools/prepare_raw.py <config.yaml> --verify   # verify raw/ only
    python tools/prepare_raw.py <config.yaml> --add      # only the listed files not yet in raw/

Configuration (YAML, to be kept OUTSIDE git and outside raw/, because it contains the real names):
    sample: ../sample                    # relative to the configuration file
    destination: my-instance/raw
    files:                               # only the listed files are copied
      - {from: "folder/file.pdf", to: "documents/file.pdf"}
    substitutions:                       # in order (longest forms first); case-insensitive
      - ["Mario Rossi", "Technician 1"]
    forbidden_extra: ["rossi"]           # searched during verification besides the substituted terms
    allowed: ["Ethernet"]                # capitalized words not proposed for review
    ocr: {command: tesseract, languages: ita+eng}   # images and scanned PDFs (Tesseract)

PDFs are REBUILT as text only, one page per original page (citations by page stay valid): writing
the replacements over the original PDF does not work (the replacement text ends up out of order
in the extraction). Layout and images are dropped; they stay in the originals.

IMAGES and PDFs WITHOUT TEXT (scans) are read with OCR: the words to redact (also names over several
words, emails, phone numbers) are covered with BLACK BOXES and the metadata (EXIF, GPS) are dropped.
OCR can misread a name, which then is not covered: the report lists the words read with low
confidence, to be checked by hand on the images before any call to the LLM.

Requires the [prep] extra (pypdfium2, to render scanned PDF pages) and Tesseract for images and scans.
"""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from curator.config import SourceRule  # noqa: E402
from curator.extractors import extract  # noqa: E402
from curator.extractors.text import decode  # noqa: E402
from curator.util import fix_ligatures, normalize_for_match  # noqa: E402

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?<!\d)(?:\+39[\s.]?)?(?:3\d{2}|0\d{1,3})[\s./-]?\d{3,4}[\s./-]?\d{3,4}(?!\d)")
CAPWORD = re.compile(r"\b[A-ZÀ-Ý][a-zà-ÿ]{2,}\b")


@dataclass
class Sub:
    term: str
    repl: str

    @property
    def rx(self) -> re.Pattern:
        return re.compile(r"(?<![\w])" + re.escape(self.term).replace(r"\ ", r"\s+") + r"(?![\w])", re.IGNORECASE)


@dataclass
class FileResult:
    dest: str
    done: Counter = field(default_factory=Counter)     # text replacements (checked as readable afterwards)
    covered: Counter = field(default_factory=Counter)  # black boxes on images, by replacement
    doubtful: list[str] = field(default_factory=list)  # words read with low OCR confidence
    notes: list[str] = field(default_factory=list)


def apply_text(text: str, subs: list[Sub], res: FileResult) -> str:
    for s in subs:
        text, n = s.rx.subn(s.repl, text)
        if n:
            res.done[s.repl] += n
    return text


# --- per format ---------------------------------------------------------------------------

def redact_text(src: Path, dst: Path, subs, res):
    text, enc = decode(src.read_bytes())
    out = apply_text(text, subs, res)
    enc = "utf-8" if enc == "utf-8-sig" else enc
    dst.write_bytes(out.encode(enc, errors="strict"))


def _docx_paragraphs(doc):
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in c.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from from_container(part)


def redact_docx(src: Path, dst: Path, subs, res):
    import docx
    d = docx.Document(str(src))
    for p in _docx_paragraphs(d):
        new = apply_text(p.text, subs, FileResult(""))
        if new != p.text:
            apply_text(p.text, subs, res)        # counting
            runs = p.runs
            if runs:
                runs[0].text = new
                for r in runs[1:]:
                    r.text = ""
    cp = d.core_properties
    for attr in ("author", "last_modified_by", "comments", "keywords", "subject", "category", "identifier"):
        setattr(cp, attr, "")
    cp.title = apply_text(cp.title or "", subs, res)
    d.save(str(dst))


def redact_xlsx(src: Path, dst: Path, subs, res):
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    for ws in wb.worksheets:
        new_title = apply_text(ws.title, subs, res)
        if new_title != ws.title:
            ws.title = new_title
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    c.value = apply_text(c.value, subs, res)
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.save(str(dst))


_PDF_CHARS = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
                            "…": "...", " ": " ", "•": "-"})


def redact_pdf(src: Path, dst: Path, subs, res):
    from pypdf import PdfReader
    from pdfwrite import text_pdf
    reader = PdfReader(str(src))
    pages = []
    for pg in reader.pages:
        text = apply_text(fix_ligatures(pg.extract_text() or ""), subs, res).translate(_PDF_CHARS)
        pages.append((text, float(pg.mediabox.width), float(pg.mediabox.height)))
    dst.write_bytes(text_pdf(pages))          # no metadata: written by hand
    res.notes.append(f"PDF rebuilt as text only ({len(reader.pages)} pages, layout and images dropped)")


# --- images: OCR + black boxes ----------------------------------------------------------------

OCR = None            # TesseractOcr set by main() from the configuration
LOW_CONF = 60.0


def _tok(s: str) -> str:
    return re.sub(r"[^\w]", "", normalize_for_match(s))


def boxes_to_cover(words, subs: list[Sub], res: FileResult) -> list[tuple[int, int, int, int]]:
    """Boxes (left, top, right, bottom) of the words to cover: substituted terms (also over several
    consecutive words of a line), emails and phone numbers."""
    by_line: dict[tuple, list] = {}
    for w in words:
        by_line.setdefault(w.line, []).append(w)
    out = []
    for line in by_line.values():
        toks = [_tok(w.text) for w in line]
        for s in subs:
            want = [t for t in (_tok(x) for x in s.term.split()) if t]
            for i in range(len(toks) - len(want) + 1):
                if want and toks[i:i + len(want)] == want:
                    ws = line[i:i + len(want)]
                    out.append((min(w.left for w in ws), min(w.top for w in ws),
                                max(w.left + w.width for w in ws), max(w.top + w.height for w in ws)))
                    res.covered[s.repl] += 1
        for w in line:
            if EMAIL.search(w.text) or PHONE.search(w.text):
                out.append((w.left, w.top, w.left + w.width, w.top + w.height))
                res.covered["[email/phone]"] += 1
            if w.conf < LOW_CONF and len(_tok(w.text)) >= 3:
                res.doubtful.append(w.text)
    return out


def _redacted_image(source, subs, res):
    """The image (path or bytes) with black boxes on the words to redact."""
    import io
    from PIL import Image, ImageDraw
    im = Image.open(io.BytesIO(source) if isinstance(source, bytes) else source).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    d = ImageDraw.Draw(im)
    for l, t, r, b in boxes_to_cover(OCR.read(buf.getvalue(), suffix=".png").words, subs, res):
        d.rectangle((l - 3, t - 3, r + 3, b + 3), fill="black")
    return im


def redact_image(src: Path, dst: Path, subs, res):
    if OCR is None:
        raise SystemExit("images need `ocr:` in the configuration (Tesseract)")
    im = _redacted_image(src, subs, res)
    fmt = "JPEG" if src.suffix.lower() in (".jpg", ".jpeg") else "PNG"
    im.save(dst, fmt, **({"quality": 95} if fmt == "JPEG" else {}))     # saved without EXIF/metadata
    res.notes.append("image: OCR + black boxes, metadata dropped")


def redact_scanned_pdf(src: Path, dst: Path, subs, res):
    """A PDF without text (scan): each page image is OCR'd, covered, and the PDF rebuilt from them."""
    import io
    import pypdfium2
    if OCR is None:
        raise SystemExit("scanned PDFs need `ocr:` in the configuration (Tesseract)")
    pages = []
    doc = pypdfium2.PdfDocument(str(src))
    try:
        for pg in doc:        # rendered: the scan may sit in nested objects that a plain extraction misses
            buf = io.BytesIO()
            pg.render(scale=200 / 72).to_pil().save(buf, "PNG")
            pages.append(_redacted_image(buf.getvalue(), subs, res))
    finally:
        doc.close()
    pages[0].save(dst, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    res.notes.append(f"scanned PDF: OCR + black boxes on {len(pages)} pages, rebuilt from the images")


def redact_any_pdf(src: Path, dst: Path, subs, res):
    from pypdf import PdfReader
    has_text = any((p.extract_text() or "").strip() for p in PdfReader(str(src)).pages)
    (redact_pdf if has_text else redact_scanned_pdf)(src, dst, subs, res)


REDACTORS = {".txt": redact_text, ".csv": redact_text, ".md": redact_text,
             ".docx": redact_docx, ".xlsx": redact_xlsx, ".pdf": redact_any_pdf,
             ".png": redact_image, ".jpg": redact_image, ".jpeg": redact_image}


# --- verification ---------------------------------------------------------------------------

def all_text(path: Path) -> str:
    """All the text the curator (or anyone opening the file) can see, metadata included."""
    ext = path.suffix.lower()
    kind = {".pdf": "pdf", ".docx": "docx", ".xlsx": "table", ".csv": "table", ".png": "image", ".jpg": "image",
            ".jpeg": "image"}.get(ext, "text")
    parts = [path.name]
    try:
        parts.append(extract(path, path.name, SourceRule("x", kind, None, True, {}), reader=OCR,
                             reader_mode="ocr" if OCR else "off").text)
    except Exception as e:  # verification must not stop on one file
        parts.append(f"[extraction failed: {e}]")
    if ext in (".docx", ".xlsx"):
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.endswith((".xml", ".rels")):
                    xml = z.read(n).decode("utf-8", "replace")
                    parts.append(re.sub(r"<[^>]+>", " ", xml))
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(path))
        parts.append(" ".join(str(v) for v in (r.metadata or {}).values() if v))
        xmp = r.trailer["/Root"].get("/Metadata")
        if xmp is not None:
            parts.append(xmp.get_object().get_data().decode("utf-8", "replace"))
    return "\n".join(parts)


def check_replacements(dest_root: Path, results: list[FileResult]) -> list[str]:
    """Every replacement made must be readable in the extracted text: otherwise the text was lost."""
    lost = []
    for r in results:
        text = all_text(dest_root / r.dest)
        for repl, n in r.done.items():
            found = text.count(repl)
            if found < n:
                lost.append(f"{r.dest}: «{repl}» replaced {n} times but readable only {found}")
    return lost


def verify(dest_root: Path, forbidden: list[str], allowed: set[str]) -> dict:
    leaks, emails, phones, words = [], Counter(), Counter(), Counter()
    norm_forbidden = [(f, normalize_for_match(f)) for f in forbidden]
    for p in sorted(dest_root.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(dest_root).as_posix()
        text = all_text(p)
        nt = normalize_for_match(text)
        squashed = re.sub(r"\s+", "", nt)
        for f, nf in norm_forbidden:
            # the space-less comparison (for "J O H N") only for long terms: "ana" is inside "banana"
            sq = re.sub(r"\s+", "", nf)
            if re.search(r"(?<!\w)" + re.escape(nf) + r"(?!\w)", nt) or (len(sq) >= 6 and sq in squashed):
                leaks.append((rel, f))
        for e in EMAIL.findall(text):
            emails[e] += 1
        for t in PHONE.findall(text):
            phones[t.strip()] += 1
        for w in CAPWORD.findall(text):
            if w not in allowed:
                words[w] += 1
    return {"leaks": leaks, "emails": emails, "phones": phones, "words": words}


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 1
    cfg_path = Path(argv[0]).resolve()
    only_verify, add = "--verify" in argv, "--add" in argv
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    global OCR
    if cfg.get("ocr"):
        from curator.ocr import TesseractOcr
        OCR = TesseractOcr(cfg["ocr"].get("command", "tesseract"), cfg["ocr"].get("languages", "eng"))
    base = cfg_path.parent
    src_root = (base / cfg["sample"]).resolve()
    dest_root = (base / cfg["destination"]).resolve()
    subs = [Sub(a, b) for a, b in cfg.get("substitutions", [])]
    forbidden = sorted({s.term for s in subs} | set(cfg.get("forbidden_extra", [])), key=str.lower)

    results: list[FileResult] = []
    if not only_verify:
        if not add and dest_root.exists() and any(p.is_file() for p in dest_root.rglob("*")):
            print(f"ERROR: {dest_root} already contains files. Empty it by hand to regenerate it, "
                  "or use --add for the files not yet there.")
            return 1
        done_by_hash: dict[str, tuple[Path, FileResult]] = {}
        for item in cfg["files"]:
            src = src_root / item["from"]
            dst = dest_root / item["to"]
            if add and dst.exists():
                continue
            if not src.is_file():
                print(f"ERROR: file not found in the sample: {item['from']}")
                return 1
            dst.parent.mkdir(parents=True, exist_ok=True)
            res = FileResult(item["to"])
            h = hashlib.sha256(src.read_bytes()).hexdigest()
            if h in done_by_hash:
                # identical copies in the sample stay identical in raw/ (same document for the curator)
                first, first_res = done_by_hash[h]
                shutil.copyfile(first, dst)
                res.done = first_res.done.copy()
                res.notes.append(f"identical copy of {first_res.dest}")
            else:
                fn = REDACTORS.get(src.suffix.lower())
                if fn is None:
                    shutil.copyfile(src, dst)
                    res.notes.append("copied without redaction (format not handled)")
                else:
                    fn(src, dst, subs, res)
                done_by_hash[h] = (dst, res)
            st = src.stat()
            os.utime(dst, (st.st_atime, st.st_mtime))   # the curator needs the date of the original file
            results.append(res)

    v = verify(dest_root, forbidden, set(cfg.get("allowed", [])))
    lost = check_replacements(dest_root, results)
    lines = ["# Redaction report", "", f"Destination: `{dest_root}`", ""]
    if results:
        lines += ["## Unreadable replacement text", ""] + ([f"- {x}" for x in lost] or ["None."]) + [""]
        lines += ["## Replacements made", "", "| file | replacements |", "|---|---|"]
        for r in results:
            done = ", ".join(f"{k} ×{n}" for k, n in sorted((r.done + r.covered).items())) or "none"
            lines.append(f"| {r.dest} | {done}{' — ' + '; '.join(r.notes) if r.notes else ''} |")
        lines.append("")
        doubtful = [r for r in results if r.doubtful]
        if doubtful:
            lines += ["## Images: words read with low OCR confidence (check them on the image)", ""]
            lines += [f"- {r.dest}: " + ", ".join(sorted(set(r.doubtful))) for r in doubtful] + [""]
    lines += ["## Verification: forbidden terms still present", ""]
    lines += [f"- **{rel}**: «{term}»" for rel, term in v["leaks"]] or ["None."]
    lines += ["", "## To review by hand", "", "Emails found:"]
    lines += [f"- {e} ×{n}" for e, n in v["emails"].most_common()] or ["- none"]
    lines += ["", "Possible phone numbers:"]
    lines += [f"- {t} ×{n}" for t, n in v["phones"].most_common()] or ["- none"]
    lines += ["", "Capitalized words (possible names), most frequent first:", ""]
    lines += [", ".join(f"{w} ×{n}" for w, n in sorted(v["words"].items(), key=lambda x: (-x[1], x[0])))]
    report = "\n".join(lines) + "\n"
    out = cfg_path.with_name(cfg_path.stem + "-report.md")
    out.write_text(report, encoding="utf-8", newline="\n")
    print(f"Files prepared: {len(results)} | forbidden terms found: {len(v['leaks'])} | "
          f"unreadable replacements: {len(lost)} | "
          f"emails: {sum(v['emails'].values())} | phones: {sum(v['phones'].values())} | "
          f"capitalized words to review: {len(v['words'])}")
    print(f"Report: {out}")
    return 2 if v["leaks"] or lost else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
